import os
import logging
from typing import List, Optional, Dict, Any
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, or_, delete
import openpyxl

from app.models.enrolled_student import EnrolledStudent
from app.schemas.enrolled_student import (
    EnrolledStudentCreate,
    EnrolledStudentUpdate,
    EnrolledStudentsStats,
)

logger = logging.getLogger("insta.services.enrolled_students")

STUDENTS_EXCEL_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "students.xlsx")
STUDENTS_SHEET_URL = "https://docs.google.com/spreadsheets/d/1B9VNXjOAmp6M6uxnsZyVuvRhqQriqgO6Dvx1dwteadE/export?format=xlsx"


class EnrolledStudentService:

    @classmethod
    async def seed_default_students(cls, session: AsyncSession, force_reload: bool = False) -> int:
        count_q = select(func.count(EnrolledStudent.id))
        count = (await session.execute(count_q)).scalar() or 0

        if count > 0 and not force_reload:
            logger.info("Enrolled students already seeded (%d records).", count)
            return count

        excel_path = STUDENTS_EXCEL_PATH
        try:
            import urllib.request
            req = urllib.request.Request(STUDENTS_SHEET_URL, headers={"User-Agent": "Mozilla/5.0"})
            data = urllib.request.urlopen(req, timeout=10).read()
            if len(data) > 1000:
                with open(STUDENTS_EXCEL_PATH, "wb") as f:
                    f.write(data)
                logger.info("Updated %s with fresh Students Google Sheets export.", STUDENTS_EXCEL_PATH)
        except Exception as e:
            logger.warning("Could not download live Students sheet (%s). Using local file %s", e, STUDENTS_EXCEL_PATH)

        if not os.path.exists(excel_path):
            logger.error("Students Excel file not found at %s", excel_path)
            return 0

        wb = openpyxl.load_workbook(excel_path, data_only=True)
        ws_proc = wb["process"]
        ws_fin = wb["finance"]

        fin_map = {}
        for r in range(2, ws_fin.max_row + 1):
            fio = str(ws_fin.cell(r, 1).value or "").strip()
            if not fio:
                continue
            fin_map[fio] = {
                "country": str(ws_fin.cell(r, 2).value or "").strip(),
                "contract_amount": str(ws_fin.cell(r, 3).value or "").strip(),
                "paid_amount": str(ws_fin.cell(r, 4).value or "0").strip(),
                "remaining_amount": str(ws_fin.cell(r, 5).value or "0").strip(),
                "seller_commission_paid": bool(ws_fin.cell(r, 6).value),
                "mentor_commission_paid": bool(ws_fin.cell(r, 7).value),
                "next_payment_condition": str(ws_fin.cell(r, 8).value or "").strip(),
                "last_payment_date": str(ws_fin.cell(r, 9).value or "").strip(),
                "finance_comment": str(ws_fin.cell(r, 10).value or "").strip(),
                "full_payment_received": bool(ws_fin.cell(r, 11).value),
            }

        if force_reload:
            await session.execute(delete(EnrolledStudent))
            await session.commit()

        created_count = 0
        for r in range(2, ws_proc.max_row + 1):
            fio = str(ws_proc.cell(r, 1).value or "").strip()
            if not fio:
                continue
            fin = fin_map.get(fio, {})
            is_done_raw = str(ws_proc.cell(r, 8).value or "").strip()
            is_done = is_done_raw.lower() in ["true", "1", "да", "баары бүттү", "бүттү"] or "отказ" in is_done_raw.lower()

            student = EnrolledStudent(
                full_name=fio,
                country=str(ws_proc.cell(r, 2).value or "").strip() or fin.get("country", ""),
                process_stage=str(ws_proc.cell(r, 3).value or "Виза").strip(),
                manager=str(ws_proc.cell(r, 4).value or "Аяна").strip(),
                next_action=str(ws_proc.cell(r, 5).value or "").strip(),
                risk=str(ws_proc.cell(r, 6).value or "").strip(),
                comment=str(ws_proc.cell(r, 7).value or "").strip(),
                is_completed=is_done,
                status_raw=is_done_raw,
                contract_amount=fin.get("contract_amount", ""),
                paid_amount=fin.get("paid_amount", "0"),
                remaining_amount=fin.get("remaining_amount", "0"),
                seller_commission_paid=fin.get("seller_commission_paid", False),
                mentor_commission_paid=fin.get("mentor_commission_paid", False),
                next_payment_condition=fin.get("next_payment_condition", ""),
                last_payment_date=fin.get("last_payment_date", ""),
                finance_comment=fin.get("finance_comment", ""),
                full_payment_received=fin.get("full_payment_received", False),
            )
            session.add(student)
            created_count += 1

        await session.commit()
        logger.info("Successfully seeded %d enrolled students into database.", created_count)
        return created_count

    @classmethod
    async def list_students(
        cls,
        session: AsyncSession,
        search: str = "",
        country: Optional[str] = None,
        manager: Optional[str] = None,
        stage: Optional[str] = None,
        risk: Optional[str] = None,
        is_completed: Optional[bool] = None,
        skip: int = 0,
        limit: int = 100,
    ) -> List[EnrolledStudent]:
        stmt = select(EnrolledStudent)

        if search:
            pattern = f"%{search.strip()}%"
            stmt = stmt.where(
                or_(
                    EnrolledStudent.full_name.ilike(pattern),
                    EnrolledStudent.country.ilike(pattern),
                    EnrolledStudent.manager.ilike(pattern),
                    EnrolledStudent.comment.ilike(pattern),
                    EnrolledStudent.next_action.ilike(pattern),
                )
            )

        if country and country != "all":
            stmt = stmt.where(EnrolledStudent.country == country)

        if manager and manager != "all":
            stmt = stmt.where(EnrolledStudent.manager == manager)

        if stage and stage != "all":
            stmt = stmt.where(EnrolledStudent.process_stage == stage)

        if risk and risk != "all":
            stmt = stmt.where(EnrolledStudent.risk == risk)

        if is_completed is not None:
            stmt = stmt.where(EnrolledStudent.is_completed == is_completed)

        stmt = stmt.order_by(EnrolledStudent.id.asc()).offset(skip).limit(limit)
        res = await session.execute(stmt)
        return list(res.scalars().all())

    @classmethod
    async def get_by_id(cls, session: AsyncSession, student_id: int) -> Optional[EnrolledStudent]:
        return await session.get(EnrolledStudent, student_id)

    @classmethod
    async def create_student(cls, session: AsyncSession, data: EnrolledStudentCreate) -> EnrolledStudent:
        student = EnrolledStudent(**data.model_dump())
        session.add(student)
        await session.commit()
        await session.refresh(student)
        return student

    @classmethod
    async def update_student(
        cls, session: AsyncSession, student_id: int, data: EnrolledStudentUpdate
    ) -> Optional[EnrolledStudent]:
        student = await session.get(EnrolledStudent, student_id)
        if not student:
            return None

        update_dict = data.model_dump(exclude_unset=True)
        for k, v in update_dict.items():
            setattr(student, k, v)

        await session.commit()
        await session.refresh(student)
        return student

    @classmethod
    async def delete_student(cls, session: AsyncSession, student_id: int) -> bool:
        student = await session.get(EnrolledStudent, student_id)
        if not student:
            return False
        await session.delete(student)
        await session.commit()
        return True

    @classmethod
    async def get_stats(cls, session: AsyncSession) -> EnrolledStudentsStats:
        total_q = select(func.count(EnrolledStudent.id))
        total = (await session.execute(total_q)).scalar() or 0

        completed_q = select(func.count(EnrolledStudent.id)).where(EnrolledStudent.is_completed == True)
        completed = (await session.execute(completed_q)).scalar() or 0

        # Countries
        c_q = select(EnrolledStudent.country, func.count(EnrolledStudent.id)).group_by(EnrolledStudent.country)
        countries_breakdown = {c: count for c, count in (await session.execute(c_q)).all() if c}

        # Stages
        s_q = select(EnrolledStudent.process_stage, func.count(EnrolledStudent.id)).group_by(EnrolledStudent.process_stage)
        stages_breakdown = {s: count for s, count in (await session.execute(s_q)).all() if s}

        # Managers
        m_q = select(EnrolledStudent.manager, func.count(EnrolledStudent.id)).group_by(EnrolledStudent.manager)
        managers_breakdown = {m: count for m, count in (await session.execute(m_q)).all() if m}

        # Risks
        r_q = select(EnrolledStudent.risk, func.count(EnrolledStudent.id)).group_by(EnrolledStudent.risk)
        risks_breakdown = {r: count for r, count in (await session.execute(r_q)).all() if r}

        return EnrolledStudentsStats(
            total_students=total,
            completed_count=completed,
            in_process_count=total - completed,
            countries_breakdown=countries_breakdown,
            stages_breakdown=stages_breakdown,
            managers_breakdown=managers_breakdown,
            risks_breakdown=risks_breakdown,
        )
