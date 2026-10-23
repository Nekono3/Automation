import os
import re
import logging
from typing import List, Optional, Tuple, Dict, Any
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, or_, update, delete
import openpyxl

from app.models.university_program import UniversityProgram
from app.schemas.university_program import (
    UniversityProgramCreate,
    UniversityProgramUpdate,
    UniversitySummary,
    ProgramsStatsResponse,
)

logger = logging.getLogger("insta.services.university_program")

SHEET_URL = "https://docs.google.com/spreadsheets/d/1wY8nJ2YbQp3TCRP6_pzpfu5iR5e_IQBvGF2_TZeiObU/export?format=xlsx"
LOCAL_EXCEL_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "turkey_universities.xlsx")

UNI_METADATA = {
    "SUBU": ("SUBU (Sakarya Uygulamalı Bilimler Üniversitesi)", "Сакарья", "Турция"),
    "SAU": ("SAU (Sakarya Üniversitesi)", "Сакарья", "Турция"),
    "ANADOLU UNIVERSITESI": ("Anadolu Üniversitesi", "Эскишехир", "Турция"),
    "ESOGU": ("ESOGU (Eskişehir Osmangazi Üniversitesi)", "Эскишехир", "Турция"),
    "BURSA TEKNIKAL UNIVERSITESI": ("Bursa Teknik Üniversitesi", "Бурса", "Турция"),
    "DUZCE UNIVERSITESI": ("Düzce Üniversitesi", "Дюздже", "Турция"),
    "АRDAHAN": ("Ardahan Üniversitesi", "Ардахан", "Турция"),
    "GIRESUN": ("Giresun Üniversitesi", "Гиресун", "Турция"),
    "BOLU": ("Bolu Abant İzzet Baysal Üniversitesi", "Болу", "Турция"),
    "KARABUK": ("Karabük Üniversitesi", "Карабюк", "Турция"),
}


class UniversityProgramService:

    @staticmethod
    def _clean(val: Any) -> str:
        if val is None:
            return ""
        return str(val).replace("\xa0", " ").strip()

    @staticmethod
    def _parse_price(p_str: str) -> Tuple[str, Optional[float], str]:
        p_str = UniversityProgramService._clean(p_str)
        if not p_str:
            return "$0", 0.0, "USD"
        curr = "USD"
        if "₺" in p_str or "TL" in p_str:
            curr = "TRY"
        elif "сом" in p_str:
            curr = "KGS"
        elif "€" in p_str or "EUR" in p_str:
            curr = "EUR"
        elif "$" in p_str:
            curr = "USD"

        cleaned = re.sub(r"[^\d,\.]", "", p_str)
        num_val = None
        if cleaned:
            if "," in cleaned and "." in cleaned:
                cleaned = cleaned.replace(".", "").replace(",", ".")
            elif "," in cleaned:
                cleaned = cleaned.replace(",", "")
            try:
                num_val = float(cleaned)
            except ValueError:
                num_val = None

        return p_str, num_val, curr

    @classmethod
    def parse_excel_file(cls, filepath: str) -> List[Dict[str, Any]]:
        """Parses the multi-sheet Excel file into structured program dictionaries."""
        wb = openpyxl.load_workbook(filepath, data_only=True)
        all_records = []

        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            uni_name, city, country = UNI_METADATA.get(sheetname, (sheetname, "Турция", "Турция"))
            current_faculty = ""
            degree_level = "Бакалавр"

            rows = []
            for r in range(1, ws.max_row + 1):
                vals = [cls._clean(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
                if any(vals):
                    rows.append(vals)

            if sheetname == "SUBU":
                for row in rows:
                    c0 = row[0] if len(row) > 0 else ""
                    c1 = row[1] if len(row) > 1 else ""
                    c2 = row[2] if len(row) > 2 else ""
                    c3 = row[3] if len(row) > 3 else ""

                    if "FAKÜLTE" in c0 or "Lisans Programları" in c1:
                        continue
                    if "MESLEK YÜKSEKOKULU" in c0 or "ÖN Lisans" in c1 or "2-х годичные" in c2:
                        degree_level = "Ön Lisans (Колледж 2 года)"
                        if c0:
                            current_faculty = c0
                        continue
                    if c0:
                        current_faculty = c0

                    prog_tr = c1
                    prog_ru = c2
                    fee = c3

                    if not prog_tr and not prog_ru:
                        continue
                    if "Контракт" in fee or "ÖĞRENİM" in c2:
                        continue

                    fee_str, fee_num, curr = cls._parse_price(fee)
                    all_records.append({
                        "university": uni_name,
                        "faculty": current_faculty or "Факультет",
                        "program_name": prog_tr or prog_ru,
                        "program_name_ru": prog_ru or prog_tr,
                        "degree_level": degree_level,
                        "tuition_fee": fee_str,
                        "tuition_fee_usd": fee_num,
                        "currency": curr,
                        "duration": "2 года" if "Ön Lisans" in degree_level else "4 года",
                        "language": "Английский" if "İNGİLİZCE" in (prog_tr + prog_ru).upper() else "Турецкий",
                        "city": city,
                        "country": country,
                        "is_active": True,
                    })

            elif sheetname == "SAU":
                for row in rows:
                    c0 = row[0] if len(row) > 0 else ""
                    c1 = row[1] if len(row) > 1 else ""
                    c2 = row[2] if len(row) > 2 else ""
                    if "Факультеты" in c0 or "Бакалавр" in c1 or "Контракт" in c2:
                        continue
                    if c0 and any(k in c0.upper() for k in ["ФАКУЛЬТЕТ", "ИНСТИТУТ", "КОЛЛЕДЖ", "ШКОЛА"]):
                        current_faculty = c0.replace("\n", " ").strip()
                    prog = c1.replace("\n", " ").strip()
                    fee = c2.strip()
                    if not prog or not fee or "Контракт" in fee:
                        continue
                    fee_str, fee_num, curr = cls._parse_price(fee)
                    lang = "Английский" if any(k in prog.upper() for k in ["АНГЛИЙСКИЙ", "ENGLISH", "İNGİLİZCE"]) else "Турецкий"
                    all_records.append({
                        "university": uni_name,
                        "faculty": current_faculty or "Факультет",
                        "program_name": prog,
                        "program_name_ru": prog,
                        "degree_level": "Бакалавр",
                        "tuition_fee": fee_str,
                        "tuition_fee_usd": fee_num,
                        "currency": curr,
                        "duration": "4 года",
                        "language": lang,
                        "city": city,
                        "country": country,
                        "is_active": True,
                    })

            elif sheetname == "ANADOLU UNIVERSITESI":
                for row in rows:
                    c0 = row[0] if len(row) > 0 else ""
                    c1 = row[1] if len(row) > 1 else ""
                    c2 = row[2] if len(row) > 2 else ""
                    if "LİSANS" in c1 or "Программы/Цена" in c0:
                        continue
                    if "FAKÜLTESİ" in c1.upper() or "ФАКУЛЬТЕТ" in c0.upper():
                        current_faculty = f"{c0} ({c1})" if c1 else c0
                        continue
                    if c2 and ("$" in c2 or any(d.isdigit() for d in c2)):
                        fee_str, fee_num, curr = cls._parse_price(c2)
                        lang = "Русский" if "РУСС" in (c0 + c1).upper() else ("Английский" if "İNGİLİZCE" in (c0 + c1).upper() else "Турецкий")
                        all_records.append({
                            "university": uni_name,
                            "faculty": current_faculty or "Факультет",
                            "program_name": c1 or c0,
                            "program_name_ru": c0 or c1,
                            "degree_level": "Бакалавр",
                            "tuition_fee": fee_str,
                            "tuition_fee_usd": fee_num,
                            "currency": curr,
                            "duration": "5 лет" if "Фармацевтика" in c0 else "4 года",
                            "language": lang,
                            "city": city,
                            "country": country,
                            "is_active": True,
                        })

            elif sheetname == "ESOGU":
                for row in rows:
                    c0 = row[0] if len(row) > 0 else ""
                    c1 = row[1] if len(row) > 1 else ""
                    c2 = row[2] if len(row) > 2 else ""
                    c3 = row[3] if len(row) > 3 else ""
                    if "Eskişehir" in c0 or "PROGRAM ADI" in c0:
                        continue
                    if "Fakültesi" in c0 or "Факультет" in c1:
                        current_faculty = f"{c0} ({c1})" if c1 else c0
                        if not c3:
                            continue
                    if c3 and ("$" in c3 or any(d.isdigit() for d in c3)):
                        fee_str, fee_num, curr = cls._parse_price(c3)
                        duration = c2 or "4 года"
                        deg = "Бакалавр"
                        if "5 лет" in duration:
                            deg = "Специалитет (5 лет)"
                        all_records.append({
                            "university": uni_name,
                            "faculty": current_faculty or "Факультет",
                            "program_name": c0,
                            "program_name_ru": c1 or c0,
                            "degree_level": deg,
                            "tuition_fee": fee_str,
                            "tuition_fee_usd": fee_num,
                            "currency": curr,
                            "duration": duration,
                            "language": "Английский" if "İngilizce" in c0 else "Турецкий",
                            "city": city,
                            "country": country,
                            "is_active": True,
                        })

            elif sheetname == "BURSA TEKNIKAL UNIVERSITESI":
                for row in rows:
                    c0 = row[0] if len(row) > 0 else ""
                    c1 = row[1] if len(row) > 1 else ""
                    c2 = row[2] if len(row) > 2 else ""
                    c3 = row[3] if len(row) > 3 else ""
                    if "FAKULTE" in c0 or "Контракт" in c3:
                        continue
                    if c0 and ("Fakültesi" in c0 or "Факультет" in c0):
                        current_faculty = c0
                    if c3 and ("$" in c3 or any(d.isdigit() for d in c3)):
                        fee_str, fee_num, curr = cls._parse_price(c3)
                        all_records.append({
                            "university": uni_name,
                            "faculty": current_faculty or "Факультет инженерии и естественных наук",
                            "program_name": c1 or c2,
                            "program_name_ru": c2 or c1,
                            "degree_level": "Бакалавр",
                            "tuition_fee": fee_str,
                            "tuition_fee_usd": fee_num,
                            "currency": curr,
                            "duration": "4 года",
                            "language": "Турецкий",
                            "city": city,
                            "country": country,
                            "is_active": True,
                        })

            elif sheetname == "DUZCE UNIVERSITESI":
                for row in rows:
                    c0 = row[0] if len(row) > 0 else ""
                    c1 = row[1] if len(row) > 1 else ""
                    c2 = row[2] if len(row) > 2 else ""
                    if "DUZCE" in c0 or "Бакалавриат" in c0 or "Programlar" in c0:
                        continue
                    if c2 and ("$" in c2 or any(d.isdigit() for d in c2)):
                        fee_str, fee_num, curr = cls._parse_price(c2)
                        lang = "Английский" if "İngilizce" in c0 or "английском" in c1 else "Турецкий"
                        duration = "5 лет" if "5 yıl" in c0 else "4 года"
                        deg = "Специалитет (5 лет)" if "5 yıl" in c0 else "Бакалавр"
                        all_records.append({
                            "university": uni_name,
                            "faculty": "Факультеты Düzce",
                            "program_name": c0,
                            "program_name_ru": c1 or c0,
                            "degree_level": deg,
                            "tuition_fee": fee_str,
                            "tuition_fee_usd": fee_num,
                            "currency": curr,
                            "duration": duration,
                            "language": lang,
                            "city": city,
                            "country": country,
                            "is_active": True,
                        })

            elif sheetname == "АRDAHAN":
                for row in rows:
                    c0 = row[0] if len(row) > 0 else ""
                    c1 = row[1] if len(row) > 1 else ""
                    c2 = row[2] if len(row) > 2 else ""
                    if "Факультеты" in c0 or "Программы" in c1:
                        continue
                    if c2 and ("$" in c2 or any(d.isdigit() for d in c2)):
                        fee_str, fee_num, curr = cls._parse_price(c2)
                        lang = "Английский" if "İngilizce" in c1 else ("Русский" if "Rus Dili" in c1 else "Турецкий")
                        all_records.append({
                            "university": uni_name,
                            "faculty": c0 or "Гуманитарный факультет",
                            "program_name": c1,
                            "program_name_ru": c1,
                            "degree_level": "Бакалавр",
                            "tuition_fee": fee_str,
                            "tuition_fee_usd": fee_num,
                            "currency": curr,
                            "duration": "4 года",
                            "language": lang,
                            "city": city,
                            "country": country,
                            "is_active": True,
                        })

            elif sheetname == "GIRESUN":
                for row in rows:
                    c0 = row[0] if len(row) > 0 else ""
                    c1 = row[1] if len(row) > 1 else ""
                    c2 = row[2] if len(row) > 2 else ""
                    c3 = row[3] if len(row) > 3 else ""
                    if "Программы" in c0 or "Годовой" in c2:
                        continue
                    if c2 and ("₺" in c2 or any(d.isdigit() for d in c2)):
                        fee_str, fee_num, curr = cls._parse_price(c2)
                        usd_val = round(fee_num / 34.0, 1) if fee_num else None
                        deg = "Бакалавр"
                        duration = "4 года"
                        if "TIP" in c0.upper():
                            duration = "6 лет"
                            deg = "Специалитет (Медицина 6 лет)"
                        elif "DİŞ" in c0.upper():
                            duration = "5 лет"
                            deg = "Специалитет (Стоматология 5 лет)"
                        all_records.append({
                            "university": uni_name,
                            "faculty": c1 or c0,
                            "program_name": c0,
                            "program_name_ru": c1 or c0,
                            "degree_level": deg,
                            "tuition_fee": f"{c2} ({c3})" if c3 else c2,
                            "tuition_fee_usd": usd_val,
                            "currency": "TRY",
                            "duration": duration,
                            "language": "Турецкий",
                            "city": city,
                            "country": country,
                            "is_active": True,
                        })

        return all_records

    @classmethod
    async def seed_default_programs(cls, session: AsyncSession, force_reload: bool = False) -> int:
        """Seeds programs from the local Excel file or Google Sheet if empty or forced."""
        count_q = select(func.count(UniversityProgram.id))
        count_res = await session.execute(count_q)
        current_count = count_res.scalar() or 0

        if current_count > 0 and not force_reload:
            logger.info("University programs already seeded (%d records).", current_count)
            return current_count

        # Try to download fresh from Google Sheets or fallback to local file
        excel_path = LOCAL_EXCEL_PATH
        try:
            import urllib.request
            logger.info("Attempting to download latest spreadsheet from Google Sheets...")
            req = urllib.request.Request(SHEET_URL, headers={"User-Agent": "Mozilla/5.0"})
            data = urllib.request.urlopen(req, timeout=10).read()
            if len(data) > 1000:
                with open(LOCAL_EXCEL_PATH, "wb") as f:
                    f.write(data)
                logger.info("Updated %s with fresh Google Sheets export (%d bytes).", LOCAL_EXCEL_PATH, len(data))
        except Exception as e:
            logger.warning("Could not download live sheet (%s). Using local file %s", e, LOCAL_EXCEL_PATH)

        if not os.path.exists(excel_path):
            logger.error("Excel file does not exist at %s", excel_path)
            return 0

        parsed_records = cls.parse_excel_file(excel_path)
        if not parsed_records:
            logger.warning("No records parsed from %s", excel_path)
            return 0

        if force_reload:
            await session.execute(delete(UniversityProgram))
            await session.commit()

        for rec in parsed_records:
            prog = UniversityProgram(**rec)
            session.add(prog)

        await session.commit()
        logger.info("Successfully seeded %d university programs into database.", len(parsed_records))
        return len(parsed_records)

    @classmethod
    async def list_programs(
        cls,
        session: AsyncSession,
        search: str = "",
        university: Optional[str] = None,
        faculty: Optional[str] = None,
        degree_level: Optional[str] = None,
        language: Optional[str] = None,
        skip: int = 0,
        limit: int = 100,
        sort_by: str = "name",
    ) -> List[UniversityProgram]:
        stmt = select(UniversityProgram).where(UniversityProgram.is_active == True)

        if search:
            search_pattern = f"%{search.strip()}%"
            stmt = stmt.where(
                or_(
                    UniversityProgram.program_name.ilike(search_pattern),
                    UniversityProgram.program_name_ru.ilike(search_pattern),
                    UniversityProgram.faculty.ilike(search_pattern),
                    UniversityProgram.university.ilike(search_pattern),
                    UniversityProgram.city.ilike(search_pattern),
                )
            )

        if university and university != "all":
            stmt = stmt.where(UniversityProgram.university == university)

        if faculty and faculty != "all":
            stmt = stmt.where(UniversityProgram.faculty == faculty)

        if degree_level and degree_level != "all":
            stmt = stmt.where(UniversityProgram.degree_level == degree_level)

        if language and language != "all":
            stmt = stmt.where(UniversityProgram.language == language)

        if sort_by == "fee_asc":
            stmt = stmt.order_by(UniversityProgram.tuition_fee_usd.asc().nulls_last())
        elif sort_by == "fee_desc":
            stmt = stmt.order_by(UniversityProgram.tuition_fee_usd.desc().nulls_last())
        elif sort_by == "university":
            stmt = stmt.order_by(UniversityProgram.university.asc(), UniversityProgram.program_name_ru.asc())
        else:
            stmt = stmt.order_by(UniversityProgram.university.asc(), UniversityProgram.faculty.asc(), UniversityProgram.program_name_ru.asc())

        stmt = stmt.offset(skip).limit(limit)
        res = await session.execute(stmt)
        return list(res.scalars().all())

    @classmethod
    async def get_by_id(cls, session: AsyncSession, program_id: int) -> Optional[UniversityProgram]:
        return await session.get(UniversityProgram, program_id)

    @classmethod
    async def create_program(cls, session: AsyncSession, data: UniversityProgramCreate) -> UniversityProgram:
        _, fee_num, curr = cls._parse_price(data.tuition_fee)
        program = UniversityProgram(
            **data.model_dump(exclude={"tuition_fee_usd", "currency"}),
            tuition_fee_usd=data.tuition_fee_usd if data.tuition_fee_usd is not None else fee_num,
            currency=data.currency or curr,
        )
        session.add(program)
        await session.commit()
        await session.refresh(program)
        return program

    @classmethod
    async def update_program(
        cls, session: AsyncSession, program_id: int, data: UniversityProgramUpdate
    ) -> Optional[UniversityProgram]:
        program = await session.get(UniversityProgram, program_id)
        if not program:
            return None

        update_dict = data.model_dump(exclude_unset=True)
        if "tuition_fee" in update_dict and "tuition_fee_usd" not in update_dict:
            _, fee_num, curr = cls._parse_price(update_dict["tuition_fee"])
            update_dict["tuition_fee_usd"] = fee_num
            if "currency" not in update_dict:
                update_dict["currency"] = curr

        for key, val in update_dict.items():
            setattr(program, key, val)

        await session.commit()
        await session.refresh(program)
        return program

    @classmethod
    async def delete_program(cls, session: AsyncSession, program_id: int) -> bool:
        program = await session.get(UniversityProgram, program_id)
        if not program:
            return False
        await session.delete(program)
        await session.commit()
        return True

    @classmethod
    async def rename_university(cls, session: AsyncSession, old_name: str, new_name: str) -> int:
        """Renames a university across all associated programs."""
        stmt = (
            update(UniversityProgram)
            .where(UniversityProgram.university == old_name)
            .values(university=new_name)
        )
        res = await session.execute(stmt)
        await session.commit()
        return res.rowcount

    @classmethod
    async def list_universities(cls, session: AsyncSession) -> List[UniversitySummary]:
        stmt = (
            select(
                UniversityProgram.university,
                UniversityProgram.city,
                UniversityProgram.country,
                func.count(UniversityProgram.id).label("count"),
            )
            .where(UniversityProgram.is_active == True)
            .group_by(UniversityProgram.university, UniversityProgram.city, UniversityProgram.country)
            .order_by(func.count(UniversityProgram.id).desc())
        )
        res = await session.execute(stmt)
        rows = res.all()

        summaries = []
        for uni, city, country, count in rows:
            deg_stmt = (
                select(func.distinct(UniversityProgram.degree_level))
                .where(UniversityProgram.university == uni, UniversityProgram.is_active == True)
            )
            deg_res = await session.execute(deg_stmt)
            degrees = [d for d in deg_res.scalars().all() if d]
            summaries.append(
                UniversitySummary(
                    name=uni,
                    city=city,
                    country=country,
                    program_count=count,
                    degrees=degrees,
                )
            )
        return summaries

    @classmethod
    async def get_stats(cls, session: AsyncSession) -> ProgramsStatsResponse:
        total_p_q = select(func.count(UniversityProgram.id)).where(UniversityProgram.is_active == True)
        total_p = (await session.execute(total_p_q)).scalar() or 0

        total_u_q = select(func.count(func.distinct(UniversityProgram.university))).where(UniversityProgram.is_active == True)
        total_u = (await session.execute(total_u_q)).scalar() or 0

        total_f_q = select(func.count(func.distinct(UniversityProgram.faculty))).where(UniversityProgram.is_active == True)
        total_f = (await session.execute(total_f_q)).scalar() or 0

        deg_q = (
            select(UniversityProgram.degree_level, func.count(UniversityProgram.id))
            .where(UniversityProgram.is_active == True)
            .group_by(UniversityProgram.degree_level)
        )
        deg_res = await session.execute(deg_q)
        degrees_breakdown = {deg: count for deg, count in deg_res.all()}

        summaries = await cls.list_universities(session)

        return ProgramsStatsResponse(
            total_programs=total_p,
            total_universities=total_u,
            total_faculties=total_f,
            degrees_breakdown=degrees_breakdown,
            universities=summaries,
        )
