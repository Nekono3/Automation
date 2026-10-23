import os
import logging
from typing import List, Optional, Dict, Any
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, delete
from sqlalchemy.orm import selectinload
import openpyxl

from app.models.seller_kpi import Seller, SellerKPIRecord
from app.schemas.seller_kpi import SellerCreate, SellerUpdate, KPIScoreAdd, SellerRead

logger = logging.getLogger("insta.services.seller_kpi")

KPI_EXCEL_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "kpi_sellers.xlsx")
KPI_SHEET_URL = "https://docs.google.com/spreadsheets/d/1Q9Mmc2dmuh-HilYdDMMM0LhUY8k52fB0EhWWD85cYjE/export?format=xlsx"


class SellerKPIService:

    @classmethod
    async def seed_default_kpi(cls, session: AsyncSession, force_reload: bool = False) -> int:
        count_q = select(func.count(Seller.id))
        count = (await session.execute(count_q)).scalar() or 0

        if count > 0 and not force_reload:
            logger.info("Sellers KPI already seeded (%d sellers).", count)
            return count

        # Try to download live sheet or fallback to local file
        excel_path = KPI_EXCEL_PATH
        try:
            import urllib.request
            req = urllib.request.Request(KPI_SHEET_URL, headers={"User-Agent": "Mozilla/5.0"})
            data = urllib.request.urlopen(req, timeout=10).read()
            if len(data) > 1000:
                with open(KPI_EXCEL_PATH, "wb") as f:
                    f.write(data)
                logger.info("Updated %s with fresh KPI Google Sheets export.", KPI_EXCEL_PATH)
        except Exception as e:
            logger.warning("Could not download live KPI sheet (%s). Using local file %s", e, KPI_EXCEL_PATH)

        if not os.path.exists(excel_path):
            logger.error("KPI Excel file not found at %s", excel_path)
            return 0

        if force_reload:
            await session.execute(delete(SellerKPIRecord))
            await session.execute(delete(Seller))
            await session.commit()

        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheetname = wb.sheetnames[0]
        ws = wb[sheetname]

        blocks = [
            ("Сезим Досумбекова", 1),
            ("Нурсулуу Сатымкулова", 9),
        ]

        total_records_created = 0
        for seller_name, start_col in blocks:
            seller = Seller(name=seller_name, role="Сатуучу (Менеджер по продажам)", is_active=True)
            session.add(seller)
            await session.flush()

            for r in range(6, ws.max_row + 1):
                period = ws.cell(r, start_col).value
                if not period or not str(period).strip():
                    continue
                period_str = str(period).strip()

                def to_float(val: Any) -> float:
                    try:
                        return float(val or 0)
                    except (ValueError, TypeError):
                        return 0.0

                punctuality = to_float(ws.cell(r, start_col + 1).value)
                calls = to_float(ws.cell(r, start_col + 2).value)
                sales = to_float(ws.cell(r, start_col + 3).value)
                meetings = to_float(ws.cell(r, start_col + 4).value)
                innovation = to_float(ws.cell(r, start_col + 5).value)
                calc_total = punctuality + calls + sales + meetings + innovation

                record = SellerKPIRecord(
                    seller_id=seller.id,
                    period=period_str,
                    punctuality=punctuality,
                    calls_meetings=calls,
                    sales=sales,
                    team_meetings=meetings,
                    innovation=innovation,
                    total_score=calc_total,
                    evaluator="РОП(Аяна эже)" if not innovation else "Арзы агай",
                )
                session.add(record)
                total_records_created += 1

        await session.commit()
        logger.info("Successfully seeded %d KPI records across sellers.", total_records_created)
        return total_records_created

    @classmethod
    async def list_sellers(cls, session: AsyncSession) -> List[SellerRead]:
        stmt = (
            select(Seller)
            .where(Seller.is_active == True)
            .options(selectinload(Seller.records))
            .order_by(Seller.id)
        )
        res = await session.execute(stmt)
        sellers = res.scalars().all()

        output = []
        for s in sellers:
            records = s.records or []
            total = sum(r.total_score for r in records)
            avg = round(total / len(records), 1) if records else 0.0
            read_obj = SellerRead.model_validate(s)
            read_obj.total_earned_score = total
            read_obj.average_score = avg
            output.append(read_obj)
        return output

    @classmethod
    async def create_seller(cls, session: AsyncSession, data: SellerCreate) -> Seller:
        seller = Seller(**data.model_dump())
        session.add(seller)
        await session.commit()
        await session.refresh(seller)
        return seller

    @classmethod
    async def delete_seller(cls, session: AsyncSession, seller_id: int) -> bool:
        seller = await session.get(Seller, seller_id)
        if not seller:
            return False
        await session.delete(seller)
        await session.commit()
        return True

    @classmethod
    async def add_or_update_score(cls, session: AsyncSession, payload: KPIScoreAdd) -> SellerKPIRecord:
        """Adds or updates score for a specific category within a period."""
        stmt = (
            select(SellerKPIRecord)
            .where(
                SellerKPIRecord.seller_id == payload.seller_id,
                SellerKPIRecord.period == payload.period.strip(),
            )
        )
        res = await session.execute(stmt)
        record = res.scalar_one_or_none()

        if not record:
            record = SellerKPIRecord(
                seller_id=payload.seller_id,
                period=payload.period.strip(),
                evaluator=payload.evaluator or "РОП",
                notes=payload.notes,
            )
            session.add(record)

        # Normalize category
        cat = payload.category.lower().strip()
        val = payload.score

        if "кечиг" in cat or "лид" in cat or "punctuality" in cat:
            record.punctuality = val
        elif "созвон" in cat or "офис" in cat or "calls" in cat:
            record.calls_meetings = val
        elif "сатуу" in cat or "продаж" in cat or "sales" in cat:
            record.sales = val
        elif "команда" in cat or "чогулуш" in cat or "meeting" in cat:
            record.team_meetings = val
        elif "жаңылык" in cat or "компания" in cat or "новост" in cat or "innovat" in cat:
            record.innovation = val
        else:
            # Default fallback to sales
            record.sales = val

        # Recalculate total score
        record.total_score = (
            record.punctuality
            + record.calls_meetings
            + record.sales
            + record.team_meetings
            + record.innovation
        )
        if payload.evaluator:
            record.evaluator = payload.evaluator
        if payload.notes:
            record.notes = payload.notes

        await session.commit()
        await session.refresh(record)
        return record

    @classmethod
    async def delete_record(cls, session: AsyncSession, record_id: int) -> bool:
        record = await session.get(SellerKPIRecord, record_id)
        if not record:
            return False
        await session.delete(record)
        await session.commit()
        return True
